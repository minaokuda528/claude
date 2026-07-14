"""PowerPoint作成ステージ.

検証・KPI・サマリーの結果から 16:9 の採用月次レポートスライドを生成する。
外部テンプレートに依存せず、python-pptx で図形を描画する。
"""

from __future__ import annotations

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

from .kpi import KpiComparison, fmt_value, fmt_mom

# ---- パレット ----
INK = RGBColor(0x11, 0x22, 0x3A)
PRIMARY = RGBColor(0x2E, 0x6F, 0xB0)
TEAL = RGBColor(0x1F, 0x9E, 0x94)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
AMBER = RGBColor(0xC7, 0x7D, 0x0A)
RED = RGBColor(0xC0, 0x3A, 0x3A)
PANEL = RGBColor(0xF4, 0xF6, 0xF9)
GRID = RGBColor(0xD9, 0xDE, 0xE6)
MUTED = RGBColor(0x6B, 0x76, 0x88)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK_SOFT = RGBColor(0x33, 0x3E, 0x52)

FONT = "IPAGothic"

EMU_W = Inches(13.333)
EMU_H = Inches(7.5)


def _no_line(shape):
    shape.line.fill.background()


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def _rect(slide, x, y, w, h, color, line=None, line_w=None, radius=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(shape_type, x, y, w, h)
    _fill(sp, color)
    if line is None:
        _no_line(sp)
    else:
        sp.line.color.rgb = line
        sp.line.width = line_w or Pt(1)
    sp.shadow.inherit = False
    if radius:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def _text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
          space_after=Pt(2), line_spacing=1.0):
    """runs: list of (text, size, color, bold) tuples -> one paragraph per tuple,
    または list of list -> 各段落が複数ランを持つ。"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for item in runs:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = align
        para.space_after = space_after
        para.space_before = Pt(0)
        para.line_spacing = line_spacing
        parts = item if isinstance(item, list) else [item]
        for (t, size, color, bold) in parts:
            r = para.add_run()
            r.text = t
            r.font.size = size
            r.font.color.rgb = color
            r.font.bold = bold
            r.font.name = FONT
    return tb


def _add_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # blank


def _header(slide, kicker, title, page_no, total):
    _rect(slide, 0, 0, EMU_W, Inches(0.14), PRIMARY)
    _text(slide, Inches(0.6), Inches(0.42), Inches(11), Inches(0.4),
          [[(kicker, Pt(12), PRIMARY, True)]])
    _text(slide, Inches(0.6), Inches(0.72), Inches(11.5), Inches(0.7),
          [[(title, Pt(26), INK, True)]])
    _rect(slide, Inches(0.6), Inches(1.42), Inches(12.13), Pt(1.5), GRID)
    _text(slide, Inches(12.0), Inches(0.5), Inches(0.9), Inches(0.3),
          [[(f"{page_no} / {total}", Pt(10), MUTED, False)]], align=PP_ALIGN.RIGHT)


# ---------------------------------------------------------------- 各スライド

def _slide_title(prs, meta, summary, total):
    slide = _add_slide(prs)
    _rect(slide, 0, 0, EMU_W, EMU_H, INK)
    _rect(slide, 0, Inches(5.0), EMU_W, Inches(0.06), PRIMARY)
    _text(slide, Inches(0.9), Inches(1.7), Inches(11.5), Inches(0.5),
          [[(meta["department"] + " ｜ 月次採用レポート", Pt(16), TEAL, True)]])
    _text(slide, Inches(0.9), Inches(2.25), Inches(11.5), Inches(1.4),
          [[(meta["period_label"], Pt(54), WHITE, True)]])
    _text(slide, Inches(0.9), Inches(3.55), Inches(11.5), Inches(0.6),
          [[(summary["headline"], Pt(18), RGBColor(0xC7, 0xD2, 0xE0), False)]])
    _text(slide, Inches(0.9), Inches(5.25), Inches(11.5), Inches(0.4),
          [[("採用データ検証 → KPI計算 → サマリー → 提言", Pt(12), MUTED, False)]])
    _text(slide, Inches(0.9), Inches(6.7), Inches(11.5), Inches(0.4),
          [[("recruitment-report-builder ｜ 自動生成レポート", Pt(10), MUTED, False)]])


def _slide_summary(prs, summary, page, total):
    slide = _add_slide(prs)
    _rect(slide, 0, 0, EMU_W, EMU_H, WHITE)
    _header(slide, "EXECUTIVE SUMMARY", "サマリー", page, total)

    top = Inches(1.75)
    col_w = Inches(5.9)
    gap = Inches(0.33)
    left_x = Inches(0.6)
    right_x = left_x + col_w + gap

    def panel(x, header_text, header_color, items, marker):
        _rect(slide, x, top, col_w, Inches(4.9), PANEL, radius=0.04)
        _rect(slide, x, top, col_w, Inches(0.55), header_color, radius=0.04)
        _text(slide, x + Inches(0.25), top + Inches(0.1), col_w - Inches(0.5), Inches(0.4),
              [[(header_text, Pt(14), WHITE, True)]])
        y = top + Inches(0.8)
        for it in items:
            _text(slide, x + Inches(0.3), y, col_w - Inches(0.6), Inches(0.7),
                  [[(marker + "  ", Pt(12), header_color, True), (it, Pt(12.5), INK_SOFT, False)]],
                  line_spacing=1.05)
            y += Inches(0.72)

    panel(left_x, "ハイライト", GREEN, summary["highlights"], "✓")
    panel(right_x, "課題", AMBER, summary["concerns"], "!")


def _slide_funnel(prs, comp, page, total):
    slide = _add_slide(prs)
    _rect(slide, 0, 0, EMU_W, EMU_H, WHITE)
    _header(slide, "RECRUITMENT FUNNEL", "採用ファネル（当月 vs 前月）", page, total)

    cur = comp.current.funnel
    prev = comp.previous.funnel
    labels = comp.current.period_label
    stage_labels = {
        "applications": "応募", "document_pass": "書類選考通過",
        "first_interview": "一次面接", "final_interview": "最終面接",
        "offers": "内定", "accepted": "内定承諾（入社）",
    }
    stages = list(cur.keys())
    max_v = max(cur.values())

    top = Inches(1.9)
    row_h = Inches(0.74)
    label_w = Inches(2.4)
    bar_x = Inches(3.15)
    bar_max_w = Inches(7.2)
    val_x = bar_x + bar_max_w + Inches(0.15)

    prev_cur = None
    for i, st in enumerate(stages):
        y = top + row_h * i
        v = cur[st]
        w = int(bar_max_w * (v / max_v)) if max_v else 0
        _text(slide, Inches(0.6), y + Inches(0.06), label_w, Inches(0.5),
              [[(stage_labels.get(st, st), Pt(13), INK, True)]])
        # 背景トラック
        _rect(slide, bar_x, y + Inches(0.05), bar_max_w, Inches(0.42), PANEL, radius=0.3)
        # 本体バー
        color = PRIMARY if st != "accepted" else GREEN
        _rect(slide, bar_x, y + Inches(0.05), max(w, Emu(1)), Inches(0.42), color, radius=0.3)
        _text(slide, val_x, y + Inches(0.05), Inches(1.4), Inches(0.42),
              [[(f"{v:,}", Pt(14), INK, True)]], anchor=MSO_ANCHOR.MIDDLE)
        # ステージ間 転換率
        if prev_cur is not None and prev_cur:
            rate = v / prev_cur
            _text(slide, val_x + Inches(1.15), y + Inches(0.05), Inches(1.9), Inches(0.42),
                  [[(f"転換 {rate*100:.0f}%", Pt(11), MUTED, False)]], anchor=MSO_ANCHOR.MIDDLE)
        prev_cur = v

    # 凡例 / MoM
    ly = top + row_h * len(stages) + Inches(0.1)
    apps_mom = (cur["applications"] - prev["applications"]) / prev["applications"]
    hire_mom = (cur["accepted"] - prev["accepted"]) / prev["accepted"]
    _rect(slide, Inches(0.6), ly, Inches(12.13), Inches(0.85), PANEL, radius=0.05)
    _text(slide, Inches(0.9), ly + Inches(0.16), Inches(11.5), Inches(0.6),
          [[("前月比  ", Pt(12), INK, True),
            (f"応募 {fmt_mom(apps_mom)}", Pt(12), GREEN if apps_mom > 0 else RED, True),
            ("     ", Pt(12), MUTED, False),
            (f"採用 {fmt_mom(hire_mom)}", Pt(12), GREEN if hire_mom > 0 else RED, True),
            (f"     全体転換率 応募→採用 {cur['accepted']/cur['applications']*100:.1f}%", Pt(12), INK_SOFT, False)]],
          anchor=MSO_ANCHOR.MIDDLE)


def _slide_scorecard(prs, comp, page, total):
    slide = _add_slide(prs)
    _rect(slide, 0, 0, EMU_W, EMU_H, WHITE)
    _header(slide, "KPI SCORECARD", "KPI達成状況（目標比・前月比）", page, total)

    cols = [
        ("KPI", Inches(3.1), PP_ALIGN.LEFT),
        ("当月", Inches(1.7), PP_ALIGN.RIGHT),
        ("前月", Inches(1.7), PP_ALIGN.RIGHT),
        ("目標", Inches(1.7), PP_ALIGN.RIGHT),
        ("前月比", Inches(1.7), PP_ALIGN.RIGHT),
        ("達成", Inches(2.43), PP_ALIGN.LEFT),
    ]
    x0 = Inches(0.6)
    top = Inches(1.85)
    header_h = Inches(0.5)
    row_h = Inches(0.66)

    # ヘッダ
    _rect(slide, x0, top, Inches(12.13), header_h, INK)
    cx = x0
    for name, w, align in cols:
        _text(slide, cx + Inches(0.15), top + Inches(0.08), w - Inches(0.3), Inches(0.35),
              [[(name, Pt(12), WHITE, True)]], align=align)
        cx += w

    y = top + header_h
    for idx, r in enumerate(comp.rows):
        bg = WHITE if idx % 2 == 0 else PANEL
        _rect(slide, x0, y, Inches(12.13), row_h, bg)
        cx = x0
        cur = fmt_value(r["current"], r["fmt"])
        prev = fmt_value(r["previous"], r["fmt"])
        tgt = fmt_value(r["target"], r["fmt"]) if r["target"] is not None else "—"
        mom = fmt_mom(r["mom"])
        mom_color = GREEN if (r["mom"] or 0) > 0 else (RED if (r["mom"] or 0) < 0 else MUTED)
        # 採用単価は下がる方が良い→色反転
        if not r["higher_is_better"] and r["mom"] is not None:
            mom_color = GREEN if r["mom"] < 0 else RED
        vals = [
            ([(r["label"], Pt(12.5), INK, True)], cols[0][2]),
            ([(cur, Pt(13), INK, True)], cols[1][2]),
            ([(prev, Pt(12), MUTED, False)], cols[2][2]),
            ([(tgt, Pt(12), INK_SOFT, False)], cols[3][2]),
            ([(mom, Pt(12), mom_color, True)], cols[4][2]),
        ]
        for (runs, align), (name, w, _a) in zip(vals, cols):
            _text(slide, cx + Inches(0.15), y + Inches(0.16), w - Inches(0.3), Inches(0.4),
                  [runs], align=align, anchor=MSO_ANCHOR.MIDDLE)
            cx += w
        # 達成バッジ
        if r["met"] is True:
            _badge(slide, cx + Inches(0.15), y + Inches(0.15), "達成", GREEN)
        elif r["met"] is False:
            _badge(slide, cx + Inches(0.15), y + Inches(0.15), "未達", AMBER)
        else:
            _text(slide, cx + Inches(0.15), y + Inches(0.16), Inches(2), Inches(0.4),
                  [[("—", Pt(12), MUTED, False)]], anchor=MSO_ANCHOR.MIDDLE)
        y += row_h

    _text(slide, Inches(0.6), y + Inches(0.18), Inches(12), Inches(0.4),
          [[("※ 採用単価は目標以下（低いほど良い）で達成と判定。数値は検証済み明細合計に基づく。",
             Pt(10), MUTED, False)]])


def _badge(slide, x, y, text, color):
    sp = _rect(slide, x, y, Inches(1.0), Inches(0.36), color, radius=0.5)
    tf = sp.text_frame
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = FONT
    sp.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE


def _slide_channels(prs, comp, page, total):
    slide = _add_slide(prs)
    _rect(slide, 0, 0, EMU_W, EMU_H, WHITE)
    _header(slide, "CHANNEL PERFORMANCE", "媒体別パフォーマンス", page, total)

    roll = comp.current.channel_rollup
    cols = [
        ("媒体", Inches(2.6), PP_ALIGN.LEFT),
        ("応募", Inches(1.5), PP_ALIGN.RIGHT),
        ("採用", Inches(1.5), PP_ALIGN.RIGHT),
        ("承諾率", Inches(1.7), PP_ALIGN.RIGHT),
        ("転換率", Inches(1.7), PP_ALIGN.RIGHT),
        ("費用", Inches(1.6), PP_ALIGN.RIGHT),
        ("採用単価", Inches(1.53), PP_ALIGN.RIGHT),
    ]
    x0 = Inches(0.6)
    top = Inches(1.9)
    header_h = Inches(0.5)
    row_h = Inches(0.6)

    _rect(slide, x0, top, Inches(12.13), header_h, INK)
    cx = x0
    for name, w, align in cols:
        _text(slide, cx + Inches(0.12), top + Inches(0.08), w - Inches(0.24), Inches(0.35),
              [[(name, Pt(12), WHITE, True)]], align=align)
        cx += w

    max_cph = max((c["cost_per_hire"] or 0) for c in roll)
    y = top + header_h
    for idx, c in enumerate(roll):
        bg = WHITE if idx % 2 == 0 else PANEL
        _rect(slide, x0, y, Inches(12.13), row_h, bg)
        cph = c["cost_per_hire"]
        cph_color = RED if cph == max_cph else INK
        cells = [
            ([(c["name"], Pt(12.5), INK, True)], cols[0][2]),
            ([(f"{c['applications']:,}", Pt(12), INK_SOFT, False)], cols[1][2]),
            ([(f"{c['hires']:,}", Pt(12.5), INK, True)], cols[2][2]),
            ([(f"{c['acceptance_rate']*100:.0f}%", Pt(12), INK_SOFT, False)], cols[3][2]),
            ([(f"{c['conversion']*100:.1f}%", Pt(12), INK_SOFT, False)], cols[4][2]),
            ([(f"¥{c['cost']:,}", Pt(12), INK_SOFT, False)], cols[5][2]),
            ([(f"¥{cph:,}" if cph else "—", Pt(12), cph_color, True)], cols[6][2]),
        ]
        cx = x0
        for (runs, align), (name, w, _a) in zip(cells, cols):
            _text(slide, cx + Inches(0.12), y + Inches(0.13), w - Inches(0.24), Inches(0.4),
                  [runs], align=align, anchor=MSO_ANCHOR.MIDDLE)
            cx += w
        y += row_h

    # 合計行
    tot_apps = sum(c["applications"] for c in roll)
    tot_hires = sum(c["hires"] for c in roll)
    tot_cost = sum(c["cost"] for c in roll)
    tot_cph = int(round(tot_cost / tot_hires)) if tot_hires else 0
    _rect(slide, x0, y, Inches(12.13), row_h, RGBColor(0xE7, 0xEC, 0xF3))
    tot_cells = [
        ([("合計 / 平均", Pt(12.5), INK, True)], cols[0][2]),
        ([(f"{tot_apps:,}", Pt(12.5), INK, True)], cols[1][2]),
        ([(f"{tot_hires:,}", Pt(12.5), INK, True)], cols[2][2]),
        ([("—", Pt(12), MUTED, False)], cols[3][2]),
        ([(f"{tot_hires/tot_apps*100:.1f}%", Pt(12.5), INK, True)], cols[4][2]),
        ([(f"¥{tot_cost:,}", Pt(12.5), INK, True)], cols[5][2]),
        ([(f"¥{tot_cph:,}", Pt(12.5), INK, True)], cols[6][2]),
    ]
    cx = x0
    for (runs, align), (name, w, _a) in zip(tot_cells, cols):
        _text(slide, cx + Inches(0.12), y + Inches(0.13), w - Inches(0.24), Inches(0.4),
              [runs], align=align, anchor=MSO_ANCHOR.MIDDLE)
        cx += w

    _text(slide, Inches(0.6), y + row_h + Inches(0.2), Inches(12), Inches(0.4),
          [[("赤字＝採用単価が最も高い媒体。リファラルは低単価・高転換で費用対効果が高い。",
             Pt(10), MUTED, False)]])


def _slide_actions(prs, summary, comp, page, total):
    slide = _add_slide(prs)
    _rect(slide, 0, 0, EMU_W, EMU_H, WHITE)
    _header(slide, "NEXT ACTIONS", "次月アクション・提言", page, total)

    top = Inches(1.9)
    y = top
    for i, a in enumerate(summary["actions"], 1):
        _rect(slide, Inches(0.6), y, Inches(0.5), Inches(0.5), PRIMARY, radius=0.2)
        _text(slide, Inches(0.6), y, Inches(0.5), Inches(0.5),
              [[(str(i), Pt(16), WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        _text(slide, Inches(1.3), y + Inches(0.05), Inches(11.3), Inches(0.6),
              [[(a, Pt(15), INK_SOFT, False)]], anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(0.82)

    # KPIターゲット再掲（帯）
    hires = next((r for r in comp.rows if r["key"] == "hires"), None)
    cph = next((r for r in comp.rows if r["key"] == "cost_per_hire"), None)
    band_y = Inches(6.15)
    _rect(slide, Inches(0.6), band_y, Inches(12.13), Inches(0.85), INK, radius=0.05)
    txt = "来月目標  "
    _text(slide, Inches(0.9), band_y + Inches(0.16), Inches(11.5), Inches(0.55),
          [[("来月目標   ", Pt(13), TEAL, True),
            (f"採用数 維持 {fmt_value(hires['current'],'int')}名+  /  ", Pt(13), WHITE, True),
            (f"採用単価 {fmt_value(cph['target'],'yen')} 以内継続", Pt(13), WHITE, True)]],
          anchor=MSO_ANCHOR.MIDDLE)


def build_presentation(meta, comp: KpiComparison, summary: dict, out_path: str) -> str:
    prs = Presentation()
    prs.slide_width = EMU_W
    prs.slide_height = EMU_H
    total = 6
    _slide_title(prs, meta, summary, total)
    _slide_summary(prs, summary, 2, total)
    _slide_funnel(prs, comp, 3, total)
    _slide_scorecard(prs, comp, 4, total)
    _slide_channels(prs, comp, 5, total)
    _slide_actions(prs, summary, comp, 6, total)
    prs.save(out_path)
    return out_path
