import sys
from pathlib import Path

import pytest
from pptx import Presentation

from conftest import DEMO, ROOT, run_script

sys.path.insert(0, str(ROOT / "scripts"))
from generate_pptx import parse_summary_sections, validate_pptx  # noqa: E402


@pytest.fixture
def report(pipeline, tmp_path):
    """デモデータから kpi・group・summary_input・pptx を生成し、パスを返す。"""
    group = tmp_path / "kpi_by_group.csv"
    r = run_script("calculate_kpi.py", "--current", pipeline["current"],
                   "--previous", pipeline["previous"], "--targets", pipeline["targets"],
                   "--output", pipeline["kpi"], "--group-output", group)
    assert r.returncode == 0, r.stderr

    sinput = tmp_path / "summary_input.json"
    r = run_script("generate_summary_input.py", "--kpi", pipeline["kpi"],
                   "--settings", DEMO / "demo-company-settings.yaml", "--output", sinput)
    assert r.returncode == 0, r.stderr

    pptx = tmp_path / "report.pptx"
    vreport = tmp_path / "pptx_validation.csv"
    r = run_script("generate_pptx.py",
                   "--kpi", pipeline["kpi"], "--group", group,
                   "--summary-input", sinput, "--summary", DEMO / "generated-summary.md",
                   "--settings", DEMO / "demo-company-settings.yaml",
                   "--output", pptx, "--validate-report", vreport)
    assert r.returncode == 0, f"生成失敗: {r.stderr}"
    return {"pptx": pptx, "sinput": sinput, "group": group, "vreport": vreport, "tmp": tmp_path}


def test_pptx_opens_and_has_10_slides(report):
    prs = Presentation(report["pptx"])  # 開けること
    assert len(prs.slides) == 10


def _all_text(prs):
    out = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        out.append(cell.text)
    return " ".join(out)


def test_key_numbers_present_and_match_csv(report):
    text = _all_text(Presentation(report["pptx"]))
    assert "98" in text            # 全体応募数
    assert "635,000" in text       # 広告費
    assert "6,480" in text         # 全体応募単価


def test_no_raw_special_values_leak(report):
    text = _all_text(Presentation(report["pptx"]))
    assert "N/A" not in text
    assert "NOT_SET" not in text
    assert "NEW" not in text
    assert "nan" not in text
    # 未設定・新規が正しく表示されている
    assert "未設定" in text


def test_validate_pptx_no_errors(report):
    import json
    sinput = json.loads(Path(report["sinput"]).read_text(encoding="utf-8"))
    findings = validate_pptx(str(report["pptx"]), 10, sinput["overall"]["metrics"])
    errors = [f for f in findings if f["level"] == "error"]
    assert not errors, f"検証エラー: {errors}"


def test_titles_are_unique(report):
    prs = Presentation(report["pptx"])
    titles = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                t = shape.text_frame.text.strip()
                if not t.isdigit():
                    titles.append(t)
                    break
    assert len(titles) == len(set(titles)), f"タイトル重複: {titles}"


def test_slide_exclusion_via_settings(pipeline, tmp_path):
    """report.slides で不要スライドを除外できる。"""
    import yaml
    settings = yaml.safe_load((DEMO / "demo-company-settings.yaml").read_text(encoding="utf-8"))
    settings["report"]["slides"] = ["cover", "main_kpi", "open_items"]
    s2 = tmp_path / "settings2.yaml"
    s2.write_text(yaml.safe_dump(settings, allow_unicode=True), encoding="utf-8")

    group = tmp_path / "g.csv"
    run_script("calculate_kpi.py", "--current", pipeline["current"],
               "--previous", pipeline["previous"], "--targets", pipeline["targets"],
               "--output", pipeline["kpi"], "--group-output", group)
    sinput = tmp_path / "si.json"
    run_script("generate_summary_input.py", "--kpi", pipeline["kpi"],
               "--settings", s2, "--output", sinput)
    pptx = tmp_path / "r2.pptx"
    r = run_script("generate_pptx.py", "--kpi", pipeline["kpi"], "--group", group,
                   "--summary-input", sinput, "--summary", DEMO / "generated-summary.md",
                   "--settings", s2, "--output", pptx)
    assert r.returncode == 0, r.stderr
    assert len(Presentation(pptx).slides) == 3


def test_make_template(tmp_path):
    out = tmp_path / "blank.pptx"
    r = run_script("generate_pptx.py", "--make-template", out)
    assert r.returncode == 0, r.stderr
    assert out.exists()
    Presentation(out)  # 開けること


def test_parse_summary_sections():
    sections = parse_summary_sections(str(DEMO / "generated-summary.md"))
    assert 1 in sections and 7 in sections
    assert "全体結果" in sections[1]["title"]
    assert sections[7]["lines"]  # 要確認に項目がある
